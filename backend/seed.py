"""Idempotent seed: parses 100cont.xlsx with openpyxl, loads questions/answers."""
import os
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from models import Question, Answer

EXCEL_PATH = os.environ.get("EXCEL_PATH", "/app/100cont.xlsx")

# Per-sheet column layout (0-indexed):
#   q_num_col, q_text_col, ans_col, pts_col
SHEET_LAYOUT = {
    "Ronda 1": ("ronda1", 2, 3, 5, 6),
    "Ronda 2": ("ronda2", 2, 3, 5, 6),
    "Final":   ("final",  2, 3, 5, 6),
    "Otras Preguntas": ("otras", 1, 2, 4, 5),
}


def _parse_sheet(ws, q_num_col, q_text_col, ans_col, pts_col):
    """Yield (order_index, q_text, [(ans_text, points), ...])."""
    rows = list(ws.iter_rows(values_only=True))
    i = 0
    while i < len(rows):
        row = rows[i]
        num = row[q_num_col] if q_num_col < len(row) else None
        txt = row[q_text_col] if q_text_col < len(row) else None
        if isinstance(num, (int, float)) and isinstance(txt, str) and txt.strip():
            order_index = int(num)
            q_text = txt.strip()
            answers = []
            j = i + 1
            while j < len(rows):
                r = rows[j]
                # Stop if next question encountered
                nxt_num = r[q_num_col] if q_num_col < len(r) else None
                nxt_txt = r[q_text_col] if q_text_col < len(r) else None
                if isinstance(nxt_num, (int, float)) and isinstance(nxt_txt, str) and nxt_txt.strip():
                    break
                a_text = r[ans_col] if ans_col < len(r) else None
                a_pts = r[pts_col] if pts_col < len(r) else None
                if isinstance(a_text, str) and a_text.strip() and isinstance(a_pts, (int, float)):
                    # Skip header rows ("Respuesta"/"Puntos")
                    if a_text.strip().lower() != "respuesta":
                        answers.append((a_text.strip(), float(a_pts)))
                j += 1
            if answers:
                # Sort desc by points, position 1..N (cap 7)
                answers.sort(key=lambda x: x[1], reverse=True)
                answers = answers[:7]
                yield order_index, q_text, answers
            i = j
        else:
            i += 1


async def seed_if_empty(session: AsyncSession):
    existing = (await session.execute(select(Question.id).limit(1))).first()
    if existing:
        print("[seed] questions already present, skipping")
        return

    if not os.path.exists(EXCEL_PATH):
        print(f"[seed] ERROR: excel not found at {EXCEL_PATH}")
        return

    print(f"[seed] loading {EXCEL_PATH}")
    wb = load_workbook(EXCEL_PATH, data_only=True)

    total = 0
    for sheet_name, layout in SHEET_LAYOUT.items():
        if sheet_name not in wb.sheetnames:
            print(f"[seed] sheet missing: {sheet_name}")
            continue
        round_key, q_num_col, q_text_col, ans_col, pts_col = layout
        ws = wb[sheet_name]
        count = 0
        for order_index, q_text, answers in _parse_sheet(ws, q_num_col, q_text_col, ans_col, pts_col):
            q = Question(round=round_key, order_index=order_index, text=q_text)
            session.add(q)
            await session.flush()
            for pos, (a_text, pts) in enumerate(answers, start=1):
                session.add(Answer(question_id=q.id, text=a_text, points=pts, position=pos))
            count += 1
        print(f"[seed] {sheet_name}: {count} questions")
        total += count

    await session.commit()
    print(f"[seed] done, total={total}")
